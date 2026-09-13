"""Generate polished, editable PPTX files from a PresentationSpec, without an LLM touching slide objects."""

import re
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

from backend.artifacts.contracts import Artifact, PresentationSpec, SlideLayout, SlideSpec, artifact_from_path
from backend.artifacts.theme import WorkbenchTheme


class PptxGenerator:
    def __init__(self, output_dir: Path):
        self.output_dir = Path(output_dir)

    def generate(self, spec: PresentationSpec, *, task_id: str) -> Artifact:
        self.output_dir.mkdir(parents=True, exist_ok=True)
        presentation = Presentation()
        presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
        presentation.core_properties.title = spec.title
        presentation.core_properties.author = spec.author or "Sovereign AI Workbench"
        for index, slide_spec in enumerate(spec.slides, start=1):
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            self._background(slide)
            self._header(slide, slide_spec.title, slide_spec.subtitle)
            self._render(slide, slide_spec)
            self._footer(slide, index)
        path = self.output_dir / f"{self._safe_name(spec.title)}.pptx"
        presentation.save(path)
        return artifact_from_path(path, artifact_type="pptx", mime_type="application/vnd.openxmlformats-officedocument.presentationml.presentation", task_id=task_id, metadata={"title": spec.title, "slide_count": len(spec.slides), "theme": WorkbenchTheme.name})

    def _render(self, slide, spec: SlideSpec) -> None:
        renderers = {
            SlideLayout.TITLE: self._title,
            SlideLayout.EXECUTIVE_SUMMARY: self._summary,
            SlideLayout.TWO_COLUMN: self._two_column,
            SlideLayout.ARCHITECTURE: self._architecture,
            SlideLayout.DATA_CHART: self._data_chart,
            SlideLayout.RECOMMENDATION: self._recommendation,
        }
        renderers[spec.layout](slide, spec)

    def _background(self, slide) -> None:
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = WorkbenchTheme.ppt_color(WorkbenchTheme.white)
        accent = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, Inches(0.18), Inches(7.5))
        accent.fill.solid(); accent.fill.fore_color.rgb = WorkbenchTheme.ppt_color(WorkbenchTheme.teal); accent.line.fill.background()

    def _header(self, slide, title: str, subtitle: str | None) -> None:
        self._text(slide, title, 0.65, 0.35, 11.9, 0.55, 25, WorkbenchTheme.navy, bold=True, font=WorkbenchTheme.title_font)
        if subtitle:
            self._text(slide, subtitle, 0.68, 0.98, 11.6, 0.3, 10.5, WorkbenchTheme.slate)
        rule = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.68), Inches(1.34), Inches(11.95), Inches(0.035))
        rule.fill.solid(); rule.fill.fore_color.rgb = WorkbenchTheme.ppt_color(WorkbenchTheme.light_blue); rule.line.fill.background()

    def _footer(self, slide, number: int) -> None:
        self._text(slide, f"SOVEREIGN AI WORKBENCH  |  {number:02d}", 0.68, 7.08, 11.8, 0.18, 8, WorkbenchTheme.slate, align=PP_ALIGN.RIGHT)

    def _title(self, slide, spec: SlideSpec) -> None:
        self._text(slide, spec.title, 0.9, 2.05, 10.5, 0.9, 34, WorkbenchTheme.navy, bold=True, font=WorkbenchTheme.title_font)
        self._text(slide, spec.subtitle or "", 0.93, 3.1, 8.8, 0.5, 18, WorkbenchTheme.slate)
        band = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(9.55), Inches(2.0), Inches(2.25), Inches(2.25))
        band.fill.solid(); band.fill.fore_color.rgb = WorkbenchTheme.ppt_color(WorkbenchTheme.light_blue); band.line.color.rgb = WorkbenchTheme.ppt_color(WorkbenchTheme.blue)
        self._text(slide, "LOCAL\n•\nTRACEABLE\n•\nSOVEREIGN", 9.8, 2.38, 1.75, 1.25, 13, WorkbenchTheme.navy, bold=True, align=PP_ALIGN.CENTER)

    def _summary(self, slide, spec: SlideSpec) -> None:
        self._bullet_box(slide, spec.bullets, 0.95, 1.75, 11.25, 4.65, 20)

    def _two_column(self, slide, spec: SlideSpec) -> None:
        self._panel(slide, spec.left_heading or "Analysis", spec.left_items or spec.bullets, 0.8, 1.65, 5.75, 4.95)
        self._panel(slide, spec.right_heading or "Implications", spec.right_items, 6.75, 1.65, 5.75, 4.95)

    def _architecture(self, slide, spec: SlideSpec) -> None:
        nodes = spec.diagram_nodes or spec.bullets
        if not nodes:
            nodes = ["Input", "Processing", "Output"]
        width = min(2.25, 10.7 / len(nodes))
        start = max(0.8, (12.45 - (len(nodes) * width + (len(nodes) - 1) * 0.35)) / 2)
        positions = []
        for index, node in enumerate(nodes):
            x = start + index * (width + 0.35)
            positions.append(x)
            shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.85), Inches(width), Inches(1.15))
            shape.fill.solid(); shape.fill.fore_color.rgb = WorkbenchTheme.ppt_color(WorkbenchTheme.light_blue)
            shape.line.color.rgb = WorkbenchTheme.ppt_color(WorkbenchTheme.blue)
            self._text(slide, node, x + 0.12, 3.14, width - 0.24, 0.5, 13, WorkbenchTheme.navy, bold=True, align=PP_ALIGN.CENTER)
            if index:
                connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(positions[index - 1] + width), Inches(3.43), Inches(x), Inches(3.43))
                connector.line.color.rgb = WorkbenchTheme.ppt_color(WorkbenchTheme.teal)

    def _data_chart(self, slide, spec: SlideSpec) -> None:
        data = ChartData(); data.categories = spec.chart_categories
        for series in spec.chart_series:
            data.add_series(series.name, series.values)
        chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.9), Inches(1.7), Inches(11.2), Inches(4.9), data).chart
        chart.has_legend = len(spec.chart_series) > 1
        if chart.has_legend:
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.value_axis.has_major_gridlines = True
        chart.category_axis.tick_labels.font.size = Pt(10)

    def _recommendation(self, slide, spec: SlideSpec) -> None:
        message = spec.recommendation or (spec.bullets[0] if spec.bullets else "Approval is requested.")
        banner = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(1.7), Inches(11.2), Inches(1.2))
        banner.fill.solid(); banner.fill.fore_color.rgb = WorkbenchTheme.ppt_color(WorkbenchTheme.navy); banner.line.fill.background()
        self._text(slide, message, 1.2, 2.02, 10.6, 0.56, 20, WorkbenchTheme.white, bold=True, align=PP_ALIGN.CENTER)
        self._bullet_box(slide, spec.decision_points or spec.bullets[1:], 1.15, 3.35, 10.6, 2.45, 17)

    def _panel(self, slide, heading: str, items: list[str], x: float, y: float, width: float, height: float) -> None:
        panel = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(width), Inches(height))
        panel.fill.solid(); panel.fill.fore_color.rgb = WorkbenchTheme.ppt_color(WorkbenchTheme.light_gray); panel.line.color.rgb = WorkbenchTheme.ppt_color(WorkbenchTheme.light_blue)
        self._text(slide, heading, x + 0.28, y + 0.27, width - 0.55, 0.35, 16, WorkbenchTheme.navy, bold=True)
        self._bullet_box(slide, items, x + 0.28, y + 0.85, width - 0.55, height - 1.1, 14)

    def _bullet_box(self, slide, items: list[str], x: float, y: float, width: float, height: float, size: float) -> None:
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
        frame = box.text_frame; frame.clear(); frame.word_wrap = True; frame.vertical_anchor = MSO_ANCHOR.TOP
        for index, item in enumerate(items):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = item; paragraph.level = 0; paragraph.font.name = WorkbenchTheme.font; paragraph.font.size = Pt(size); paragraph.font.color.rgb = WorkbenchTheme.ppt_color(WorkbenchTheme.dark); paragraph.space_after = Pt(12)
            paragraph.text = "• " + paragraph.text

    def _text(self, slide, value: str, x: float, y: float, width: float, height: float, size: float, color: str, *, bold: bool = False, font: str | None = None, align=None) -> None:
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
        paragraph = box.text_frame.paragraphs[0]; paragraph.text = value; paragraph.font.name = font or WorkbenchTheme.font; paragraph.font.size = Pt(size); paragraph.font.bold = bold; paragraph.font.color.rgb = WorkbenchTheme.ppt_color(color)
        if align is not None: paragraph.alignment = align

    @staticmethod
    def _safe_name(value: str) -> str:
        return re.sub(r"[^A-Za-z0-9._-]+", "-", value).strip("-.")[:80] or "presentation"
