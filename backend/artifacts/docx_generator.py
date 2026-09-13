"""Local-only generator for structured approval notes."""

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from backend.artifacts.contracts import ApprovalNoteSpec, Artifact, artifact_from_path
from backend.artifacts.theme import WorkbenchTheme


class DocxGenerator:
    def __init__(self, output_dir: Path):
        self.output_dir = Path(output_dir)

    def generate(self, spec: ApprovalNoteSpec, *, task_id: str) -> Artifact:
        self.output_dir.mkdir(parents=True, exist_ok=True)
        document = Document()
        section = document.sections[0]
        section.top_margin = section.bottom_margin = Inches(0.65)
        section.left_margin = section.right_margin = Inches(0.8)
        self._set_defaults(document)

        heading = document.add_heading(spec.title, level=0)
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._color_first_run(heading, WorkbenchTheme.navy)
        if spec.reference_number:
            paragraph = document.add_paragraph(f"Reference: {spec.reference_number}")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        document.add_paragraph(f"Date: {spec.document_date.strftime('%d %B %Y')}")
        if spec.recipient:
            document.add_paragraph(f"To: {spec.recipient}")
        subject = document.add_paragraph()
        subject.add_run("Subject: ").bold = True
        subject.add_run(spec.subject)

        self._section(document, "Purpose", spec.purpose)
        if spec.background:
            self._section(document, "Background", spec.background)
        if spec.findings:
            document.add_heading("Key Findings", level=1)
            for finding in spec.findings:
                document.add_paragraph(finding, style="List Bullet")
        self._section(document, "Recommendation", spec.recommendation)
        self._section(document, "Approval Requested", spec.requested_approval)
        if spec.source_references:
            document.add_heading("Source References", level=1)
            for reference in spec.source_references:
                document.add_paragraph(reference, style="List Bullet")
        if spec.prepared_by:
            document.add_paragraph().add_run(f"Prepared by: {spec.prepared_by}").italic = True

        filename = f"{self._safe_name(spec.subject)}.docx"
        path = self.output_dir / filename
        document.save(path)
        return artifact_from_path(path, artifact_type="docx", mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", task_id=task_id, metadata={"subject": spec.subject, "theme": WorkbenchTheme.name})

    @staticmethod
    def _safe_name(value: str) -> str:
        return re.sub(r"[^A-Za-z0-9._-]+", "-", value).strip("-.")[:80] or "approval-note"

    @staticmethod
    def _color_first_run(paragraph, color: str) -> None:
        if paragraph.runs:
            paragraph.runs[0].font.color.rgb = WorkbenchTheme.docx_color(color)

    @staticmethod
    def _section(document: Document, heading: str, text: str) -> None:
        document.add_heading(heading, level=1)
        document.add_paragraph(text)

    @staticmethod
    def _set_defaults(document: Document) -> None:
        normal = document.styles["Normal"]
        normal.font.name = WorkbenchTheme.font
        normal.font.size = Pt(10.5)
        for name in ("Heading 1", "Heading 2"):
            style = document.styles[name]
            style.font.name = WorkbenchTheme.title_font
            style.font.color.rgb = WorkbenchTheme.docx_color(WorkbenchTheme.navy)
