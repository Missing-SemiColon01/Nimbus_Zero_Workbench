"""Unit tests for OfficeArtifactValidator."""

from pathlib import Path
import pytest
from docx import Document
from pptx import Presentation
import openpyxl

from backend.artifacts.contracts import Artifact
from backend.artifacts.office_validator import OfficeArtifactValidator


def test_validate_docx_with_paragraphs_only(tmp_path: Path):
    doc_path = tmp_path / "sample.docx"
    doc = Document()
    doc.add_heading("Safety Report", level=1)
    doc.add_paragraph("All equipment passed inspection successfully.")
    doc.save(str(doc_path))

    artifact = Artifact(
        id="art-1",
        task_id="task-1",
        filename="sample.docx",
        type="docx",
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        storage_uri=str(doc_path),
        size_bytes=doc_path.stat().st_size,
    )

    validator = OfficeArtifactValidator()
    result = validator.validate(artifact, required_text=["Safety Report", "equipment passed"])

    assert result.valid is True
    assert result.checks["artifact_opened"] is True
    assert result.checks["has_content"] is True
    assert result.checks["required_text_present"] is True
    assert len(result.findings) == 0


def test_validate_docx_with_tables_only(tmp_path: Path):
    """Verify that docx files structured entirely with tables pass validation."""
    doc_path = tmp_path / "table_doc.docx"
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Item"
    table.cell(0, 1).text = "Status"
    table.cell(1, 0).text = "Compressor"
    table.cell(1, 1).text = "Operational"
    doc.save(str(doc_path))

    artifact = Artifact(
        id="art-2",
        task_id="task-2",
        filename="table_doc.docx",
        type="docx",
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        storage_uri=str(doc_path),
        size_bytes=doc_path.stat().st_size,
    )

    validator = OfficeArtifactValidator()
    result = validator.validate(artifact, required_text=["Compressor", "Operational"])

    assert result.valid is True
    assert result.checks["artifact_opened"] is True
    assert result.checks["has_content"] is True
    assert result.checks["required_text_present"] is True
    assert len(result.findings) == 0


def test_validate_docx_with_paragraphs_and_tables(tmp_path: Path):
    doc_path = tmp_path / "mixed_doc.docx"
    doc = Document()
    doc.add_heading("Executive Summary", level=1)
    doc.add_paragraph("Overview of system metrics:")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Pressure"
    table.cell(1, 1).text = "105 PSI"
    doc.save(str(doc_path))

    artifact = Artifact(
        id="art-3",
        task_id="task-3",
        filename="mixed_doc.docx",
        type="docx",
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        storage_uri=str(doc_path),
        size_bytes=doc_path.stat().st_size,
    )

    validator = OfficeArtifactValidator()
    result = validator.validate(artifact, required_text=["Executive Summary", "105 PSI"])

    assert result.valid is True
    assert result.checks["has_content"] is True
    assert result.checks["required_text_present"] is True


def test_validate_pptx(tmp_path: Path):
    pptx_path = tmp_path / "slides.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Quarterly Business Review"
    prs.save(str(pptx_path))

    artifact = Artifact(
        id="art-4",
        task_id="task-4",
        filename="slides.pptx",
        type="pptx",
        mime_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        storage_uri=str(pptx_path),
        size_bytes=pptx_path.stat().st_size,
    )

    validator = OfficeArtifactValidator()
    result = validator.validate(artifact, required_text=["Quarterly Business Review"])

    assert result.valid is True
    assert result.checks["artifact_opened"] is True
    assert result.checks["has_content"] is True
    assert result.checks["required_text_present"] is True


def test_validate_xlsx(tmp_path: Path):
    xlsx_path = tmp_path / "sheet.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Telemetry"
    ws["A1"] = "Timestamp"
    ws["B1"] = "Temperature"
    ws["A2"] = "2026-09-20"
    ws["B2"] = 72.5
    wb.save(str(xlsx_path))

    artifact = Artifact(
        id="art-5",
        task_id="task-5",
        filename="sheet.xlsx",
        type="xlsx",
        mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        storage_uri=str(xlsx_path),
        size_bytes=xlsx_path.stat().st_size,
    )

    validator = OfficeArtifactValidator()
    result = validator.validate(artifact, required_text=["Telemetry", "Temperature", "72.5"])

    assert result.valid is True
    assert result.checks["artifact_opened"] is True
    assert result.checks["has_content"] is True
    assert result.checks["required_text_present"] is True


def test_validate_unsupported_type(tmp_path: Path):
    txt_path = tmp_path / "test.txt"
    txt_path.write_text("plain text")

    artifact = Artifact(
        id="art-6",
        task_id="task-6",
        filename="test.txt",
        type="txt",
        mime_type="text/plain",
        storage_uri=str(txt_path),
        size_bytes=len("plain text"),
    )

    validator = OfficeArtifactValidator()
    result = validator.validate(artifact)

    assert result.valid is False
    assert result.checks["supported_type"] is False
    assert any("must be docx, pptx, or xlsx" in f for f in result.findings)


def test_validate_nonexistent_file(tmp_path: Path):
    missing_path = tmp_path / "missing.docx"

    artifact = Artifact(
        id="art-7",
        task_id="task-7",
        filename="missing.docx",
        type="docx",
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        storage_uri=str(missing_path),
        size_bytes=0,
    )

    validator = OfficeArtifactValidator()
    result = validator.validate(artifact)

    assert result.valid is False
    assert result.checks["file_exists"] is False
    assert any("does not exist" in f for f in result.findings)


def test_validate_empty_docx(tmp_path: Path):
    empty_path = tmp_path / "empty.docx"
    doc = Document()
    doc.save(str(empty_path))

    artifact = Artifact(
        id="art-8",
        task_id="task-8",
        filename="empty.docx",
        type="docx",
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        storage_uri=str(empty_path),
        size_bytes=empty_path.stat().st_size,
    )

    validator = OfficeArtifactValidator()
    result = validator.validate(artifact)

    assert result.valid is False
    assert result.checks["has_content"] is False
    assert any("no readable content" in f for f in result.findings)


def test_validate_missing_required_text(tmp_path: Path):
    doc_path = tmp_path / "sample.docx"
    doc = Document()
    doc.add_paragraph("Actual content here.")
    doc.save(str(doc_path))

    artifact = Artifact(
        id="art-9",
        task_id="task-9",
        filename="sample.docx",
        type="docx",
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        storage_uri=str(doc_path),
        size_bytes=doc_path.stat().st_size,
    )

    validator = OfficeArtifactValidator()
    result = validator.validate(artifact, required_text=["Expected Keyword Missing"])

    assert result.valid is False
    assert result.checks["required_text_present"] is False
    assert any("missing expected content" in f for f in result.findings)


def test_validate_corrupted_file(tmp_path: Path):
    corrupt_path = tmp_path / "corrupt.docx"
    corrupt_path.write_bytes(b"not a valid zip file")

    artifact = Artifact(
        id="art-10",
        task_id="task-10",
        filename="corrupt.docx",
        type="docx",
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        storage_uri=str(corrupt_path),
        size_bytes=len(b"not a valid zip file"),
    )

    validator = OfficeArtifactValidator()
    result = validator.validate(artifact)

    assert result.valid is False
    assert result.checks["artifact_opened"] is False
    assert any("could not be opened" in f for f in result.findings)

