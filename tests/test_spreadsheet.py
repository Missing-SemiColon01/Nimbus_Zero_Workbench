"""
tests/test_spreadsheet.py
=========================
Tests for local XLSX spreadsheet generator, tool adapter, and validation:
- SpreadsheetSpec and SheetSpec validation
- XlsxGenerator generating structured .xlsx with multiple sheets and summary row
- SpreadsheetCreateTool executing within task context and adhering to require_approval policy
- OfficeArtifactValidator and ArtifactValidateTool validating generated .xlsx files
"""

from pathlib import Path
import asyncio
from openpyxl import load_workbook

from backend.artifacts.contracts import ColumnDef, SheetSpec, SpreadsheetSpec
from backend.artifacts.xlsx_generator import XlsxGenerator
from backend.artifacts.tools import SpreadsheetCreateTool
from backend.artifacts.validation_tool import ArtifactValidateTool
from backend.security.policy import PolicyDecision, PolicyEngine


def test_xlsx_generator_creates_workbook(tmp_path: Path):
    spec = SpreadsheetSpec(
        title="Quarterly Equipment Inspection Log",
        sheets=[
            SheetSpec(
                title="Pressure Valves",
                columns=[
                    ColumnDef(key="tag", header="Valve Tag"),
                    ColumnDef(key="pressure_psi", header="Operating PSI"),
                    ColumnDef(key="status", header="Status"),
                ],
                rows=[
                    {"tag": "SV-401", "pressure_psi": 120.5, "status": "PASS"},
                    {"tag": "SV-402", "pressure_psi": 135.0, "status": "PASS"},
                    {"tag": "SV-403", "pressure_psi": 149.2, "status": "PASS"},
                ],
                summary_row=True,
            ),
            SheetSpec(
                title="Pump Inventory",
                columns=[
                    ColumnDef(key="id", header="Pump ID"),
                    ColumnDef(key="rpm", header="RPM"),
                ],
                rows=[
                    {"id": "P-101", "rpm": 1750},
                    {"id": "P-102", "rpm": 3450},
                ],
            ),
        ],
    )

    generator = XlsxGenerator(tmp_path)
    artifact = generator.generate(spec, task_id="task-sheet-001")

    assert artifact.type == "xlsx"
    assert artifact.mime_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert Path(artifact.storage_uri).exists()

    # Verify content using openpyxl directly
    wb = load_workbook(artifact.storage_uri)
    assert len(wb.sheetnames) == 2
    assert "Pressure Valves" in wb.sheetnames
    assert "Pump Inventory" in wb.sheetnames

    ws1 = wb["Pressure Valves"]
    assert ws1.cell(row=1, column=1).value == "Quarterly Equipment Inspection Log"
    assert ws1.cell(row=4, column=1).value == "Valve Tag"
    assert ws1.cell(row=5, column=1).value == "SV-401"
    # Row 8 should be summary row
    assert ws1.cell(row=8, column=1).value == "Total"
    assert "=SUM(B5:B7)" in str(ws1.cell(row=8, column=2).value)
    wb.close()


def test_spreadsheet_tool_requires_approval_and_task_id(tmp_path: Path):
    policy = PolicyEngine()
    assert policy.evaluate("spreadsheet.create") == PolicyDecision.REQUIRE_APPROVAL

    tool = SpreadsheetCreateTool(tmp_path)
    # Fail without task_id
    res_no_task = asyncio.run(tool.execute({"title": "Test", "sheets": []}, {}))
    assert not res_no_task.success
    assert "context.task_id" in res_no_task.error

    # Succeed with task_id
    payload = {
        "title": "Inspection Metrics",
        "sheets": [
            {
                "title": "Metrics",
                "columns": [
                    {"key": "metric", "header": "Metric"},
                    {"key": "value", "header": "Value"},
                ],
                "rows": [
                    {"metric": "Temperature", "value": 72.5},
                ],
            }
        ],
    }
    res = asyncio.run(tool.execute(payload, {"task_id": "task-tool-001"}))
    assert res.success
    assert res.output["type"] == "xlsx"
    assert Path(res.artifacts[0]).exists()


def test_validate_tool_validates_created_xlsx(tmp_path: Path):
    artifacts_dir = tmp_path / "artifacts"
    previews_dir = tmp_path / "previews"
    artifacts_dir.mkdir(parents=True, exist_ok=True)
    previews_dir.mkdir(parents=True, exist_ok=True)

    tool = SpreadsheetCreateTool(artifacts_dir)
    payload = {
        "title": "Audit Summary",
        "sheets": [
            {
                "title": "Findings",
                "columns": [
                    {"key": "item", "header": "Audit Item"},
                    {"key": "rating", "header": "Score"},
                ],
                "rows": [
                    {"item": "Pipe wall thickness", "rating": 9.5},
                    {"item": "Flange alignment", "rating": 8.0},
                ],
            }
        ],
    }
    create_res = asyncio.run(tool.execute(payload, {"task_id": "task-audit-001"}))
    assert create_res.success
    artifact = create_res.output

    val_tool = ArtifactValidateTool(previews_dir)
    val_res = asyncio.run(
        val_tool.execute(
            {
                "artifact": artifact,
                "required_text": ["Audit Summary", "Pipe wall thickness"],
                "render_preview": False,
            },
            {"task_id": "task-audit-001"},
        )
    )

    assert val_res.success
    assert val_res.output["valid"] is True
    assert val_res.output["checks"]["has_content"] is True
    assert val_res.output["checks"]["required_text_present"] is True

