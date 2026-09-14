from pathlib import Path
from types import SimpleNamespace

import fitz
from docx import Document
from pptx import Presentation

from backend.artifacts import (
    ApprovalNoteSpec,
    DocxGenerator,
    PdfApprovalNoteGenerator,
    PdfArtifactValidator,
    OfficeArtifactValidator,
    LibreOfficeRenderer,
    PptxGenerator,
    PresentationSpec,
    SlideLayout,
    SlideSpec,
)
from backend.artifacts.contracts import ChartSeries


def test_docx_generator_creates_openable_approval_note(tmp_path: Path):
    artifact = DocxGenerator(tmp_path).generate(
        ApprovalNoteSpec(
            subject="Pump P-101 repair approval",
            recipient="Maintenance Manager",
            purpose="Obtain approval for the proposed repair.",
            findings=["Abnormal vibration observed during inspection."],
            recommendation="Repair the bearing assembly during the next shutdown.",
            requested_approval="Approve the repair work order.",
            source_references=["Inspection report IR-101, page 3"],
        ),
        task_id="task-123",
    )

    path = Path(artifact.storage_uri)
    assert path.exists()
    assert artifact.type == "docx"
    assert "Pump P-101 repair approval" in "\n".join(p.text for p in Document(path).paragraphs)


def test_pptx_generator_supports_all_day_one_layouts(tmp_path: Path):
    slides = [
        SlideSpec(layout=SlideLayout.TITLE, title="Inspection review", subtitle="September 2026"),
        SlideSpec(layout=SlideLayout.EXECUTIVE_SUMMARY, title="Executive summary", bullets=["Action is required"]),
        SlideSpec(layout=SlideLayout.TWO_COLUMN, title="Assessment", left_heading="Findings", left_items=["Vibration"], right_heading="Impact", right_items=["Reliability risk"]),
        SlideSpec(layout=SlideLayout.ARCHITECTURE, title="Workflow", diagram_nodes=["Report", "Analysis", "Approval"]),
        SlideSpec(layout=SlideLayout.DATA_CHART, title="Trend", chart_categories=["Jan", "Feb"], chart_series=[ChartSeries(name="Risk", values=[2, 4])]),
        SlideSpec(layout=SlideLayout.RECOMMENDATION, title="Decision", recommendation="Approve the repair", decision_points=["Assign owner", "Schedule shutdown"]),
    ]
    artifact = PptxGenerator(tmp_path).generate(PresentationSpec(title="Inspection review", slides=slides), task_id="task-123")

    path = Path(artifact.storage_uri)
    assert path.exists()
    assert artifact.type == "pptx"
    assert artifact.metadata["slide_count"] == 6
    assert len(Presentation(path).slides) == 6


def test_pdf_generator_creates_readable_approval_note(tmp_path: Path):
    artifact = PdfApprovalNoteGenerator(tmp_path).generate(
        ApprovalNoteSpec(
            subject="Pump P-101 repair approval",
            purpose="Obtain approval for the proposed repair.",
            findings=["Abnormal vibration observed during inspection."],
            recommendation="Repair the bearing assembly during the next shutdown.",
            requested_approval="Approve the repair work order.",
        ),
        task_id="task-123",
    )

    document = fitz.open(artifact.storage_uri)
    assert document.page_count == 1
    assert "Pump P-101 repair approval" in document[0].get_text()
    document.close()


def test_pdf_validator_checks_content_and_creates_local_preview(tmp_path: Path):
    artifact = PdfApprovalNoteGenerator(tmp_path).generate(
        ApprovalNoteSpec(
            subject="Pump P-101 repair approval",
            purpose="Obtain approval for the proposed repair.",
            recommendation="Repair the bearing assembly during the next shutdown.",
            requested_approval="Approve the repair work order.",
        ),
        task_id="task-123",
    )

    validation = PdfArtifactValidator(tmp_path / "previews").validate(
        artifact,
        required_text=["Pump P-101 repair approval", "Approval Requested"],
    )

    assert validation.valid
    assert all(validation.checks.values())
    assert validation.preview_path
    assert Path(validation.preview_path).exists()


def test_office_validator_checks_docx_and_pptx_content(tmp_path: Path):
    docx = DocxGenerator(tmp_path).generate(
        ApprovalNoteSpec(
            subject="Pump P-101 repair approval",
            purpose="Obtain approval for the proposed repair.",
            recommendation="Repair the bearing assembly.",
            requested_approval="Approve the repair work order.",
        ),
        task_id="task-123",
    )
    pptx = PptxGenerator(tmp_path).generate(
        PresentationSpec(
            title="Inspection review",
            slides=[SlideSpec(layout=SlideLayout.EXECUTIVE_SUMMARY, title="Executive summary", bullets=["Action is required"])],
        ),
        task_id="task-123",
    )
    validator = OfficeArtifactValidator()

    assert validator.validate(docx, required_text=["Pump P-101 repair approval", "Approval Requested"]).valid
    assert validator.validate(pptx, required_text=["Executive summary", "Action is required"]).valid


def test_libreoffice_renderer_converts_and_previews_docx(tmp_path: Path, monkeypatch):
    artifact = DocxGenerator(tmp_path).generate(
        ApprovalNoteSpec(subject="Pump approval", purpose="Purpose", recommendation="Repair.", requested_approval="Approve."),
        task_id="task-123",
    )
    soffice = tmp_path / "soffice.exe"
    soffice.touch()

    def fake_run(command, **_kwargs):
        output_dir = Path(command[command.index("--outdir") + 1])
        pdf = fitz.open()
        page = pdf.new_page()
        page.insert_text((72, 72), "Rendered approval note")
        pdf.save(output_dir / "Pump-approval.pdf")
        pdf.close()
        return SimpleNamespace(returncode=0, stdout="", stderr="")

    monkeypatch.setattr("backend.artifacts.office_renderer.subprocess.run", fake_run)
    renderer = LibreOfficeRenderer(tmp_path / "renders", soffice_path=soffice)
    pdf_path, preview_path = renderer.render(Path(artifact.storage_uri), preview_name=artifact.id)

    assert pdf_path.exists()
    assert preview_path.exists()
