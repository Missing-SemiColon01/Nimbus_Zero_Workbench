"""
pptx_helpers.py - Flexible presentation design SDK for LLM-written python-pptx code.

The LLM has full visual and design freedom:
- Layout composition, x/y coordinates, dimensions, margins, and padding.
- Colors, themes, typography, font families, sizes, weights, and alignments.
- Rendering primitives: slides, text boxes, shapes, lines, images, tables, charts, cards, KPIs, timelines.
- Zero fixed slide limits: create 1, 3, 7, 15, or however many slides are needed.

Safety and python-pptx intricacies are handled inside the helper primitives so the LLM
never writes raw XML or low-level COM/DrawingML code.

Example:
    from pptx_helpers import Deck, split, grid, pad

    d = Deck(theme="teal")
    # Custom blank slide with custom color
    s1 = d.add_slide(bg="0B192C")
    d.textbox(s1, "Executive Strategic Review", box=(1.0, 1.2, 11.333, 1.2), size=40, color="FFFFFF", bold=True)
    d.textbox(s1, "Q3 Operations, Infrastructure & Growth Analysis", box=(1.0, 2.4, 11.333, 0.6), size=18, color="00A6A6")
    
    # Custom 3-column metric cards
    left, mid, right = split((1.0, 3.6, 11.333, 2.8), [1, 1, 1])
    for box, (val, lab) in zip([left, mid, right], [("$12.4M", "Quarterly ARR"), ("99.98%", "System Uptime"), ("-32%", "Support Latency")]):
        d.shape(s1, kind="rounded_rect", box=box, fill="1E3E62", radius=0.08)
        d.textbox(s1, val, box=pad(box, 0.3), size=36, color="00A6A6", bold=True, align="center")
        d.textbox(s1, lab, box=(box[0] + 0.3, box[1] + 1.6, box[2] - 0.6, 0.6), size=15, color="E0E0E0", align="center")

    d.save("deck.pptx")
"""
from __future__ import annotations

import math
from pathlib import Path
from typing import Any, Sequence

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

# --------------------------------------------------------------------------- #
# Canvas Geometry & Constants
# --------------------------------------------------------------------------- #
SW, SH = 13.333, 7.5         # 16:9 widescreen canvas (inches)
M = 0.7                      # Standard side margin
GAP = 0.35                   # Standard element gap
CW = SW - 2 * M              # Content width (11.933 in)
CH = SH - 2 * M              # Content height (6.1 in)
BODY = (M, 1.95, CW, 4.85)   # Default body area below typical headers


# --------------------------------------------------------------------------- #
# Layout & Spacing Helpers
# --------------------------------------------------------------------------- #
def split(box: tuple[float, float, float, float], ratios: Sequence[float], gap: float = GAP) -> list[tuple[float, float, float, float]]:
    """Split a box horizontally into side-by-side boxes by ratios, e.g. split(BODY, [2, 1])."""
    x, y, w, h = box
    avail = w - gap * (len(ratios) - 1)
    total = float(sum(ratios))
    out, cx = [], x
    for r in ratios:
        cw = avail * r / total
        out.append((cx, y, cw, h))
        cx += cw + gap
    return out


def vsplit(box: tuple[float, float, float, float], ratios: Sequence[float], gap: float = GAP) -> list[tuple[float, float, float, float]]:
    """Split a box vertically into stacked boxes by ratios, e.g. vsplit(BODY, [1, 2])."""
    x, y, w, h = box
    avail = h - gap * (len(ratios) - 1)
    total = float(sum(ratios))
    out, cy = [], y
    for r in ratios:
        rh = avail * r / total
        out.append((x, cy, w, rh))
        cy += rh + gap
    return out


def grid(box: tuple[float, float, float, float], cols: int, rows: int = 1, gap: float = GAP) -> list[tuple[float, float, float, float]]:
    """Row-major grid of equal cells: grid(box, cols=3, rows=2)."""
    x, y, w, h = box
    cw = (w - gap * (cols - 1)) / cols
    rh = (h - gap * (rows - 1)) / rows
    return [(x + c * (cw + gap), y + r * (rh + gap), cw, rh)
            for r in range(rows) for c in range(cols)]


def pad(box: tuple[float, float, float, float], p: float) -> tuple[float, float, float, float]:
    """Inset a box by padding p on all sides: pad((x, y, w, h), 0.2)."""
    x, y, w, h = box
    return (x + p, y + p, max(0.1, w - 2 * p), max(0.1, h - 2 * p))


# --------------------------------------------------------------------------- #
# Text measurement & typography utilities
# --------------------------------------------------------------------------- #
_WIDE_FONTS = {"Cambria", "Georgia", "Times New Roman", "Bookman Old Style", "Century Schoolbook"}


def est_height(text: str, w: float, pt: float, font: str = "Calibri", bold: bool = False, para_gap_pt: float = 0) -> float:
    """Estimated rendered height (inches) of text in a box w inches wide."""
    k = 0.52 if font in _WIDE_FONTS else 0.49
    if bold:
        k *= 1.06
    cpl = max(1, int(w * 72 / (pt * k) * 0.94))
    paras = str(text).split("\n")
    lines = sum(max(1, math.ceil(len(p) / cpl)) for p in paras)
    return (lines * pt * 1.2 + para_gap_pt * (len(paras) - 1)) / 72


def fit_pt(text: str, w: float, h: float, pt: float, min_pt: float = 10, font: str = "Calibri", bold: bool = False, para_gap_pt: float = 0) -> float:
    """Find maximum font size <= pt that fits within height h."""
    while pt > min_pt and est_height(text, w, pt, font, bold, para_gap_pt) > h:
        pt -= 1
    return pt


# --------------------------------------------------------------------------- #
# Color & Theme Primitives
# --------------------------------------------------------------------------- #
def _mix(a: str, b: str, t: float) -> str:
    """Blend hex colour a toward b by fraction t (0..1)."""
    ca = [int(a[i:i + 2], 16) for i in (0, 2, 4)]
    cb = [int(b[i:i + 2], 16) for i in (0, 2, 4)]
    return "".join(f"{round(x + (y - x) * t):02X}" for x, y in zip(ca, cb))


def _rgb(val: str) -> RGBColor:
    """Normalize 6-digit hex string (with or without #) to RGBColor."""
    clean = val.lstrip("#")
    if len(clean) != 6:
        clean = "000000"
    return RGBColor.from_string(clean)


THEMES: dict[str, dict[str, Any]] = {
    "midnight": dict(dark="1E2761", primary="1E2761", accent="3B6FE8", on_accent="FFFFFF",
                     accent_text="2F5BC4", ink="1A1F3D", muted="5A6483", sub="CADCFC", bg="FFFFFF",
                     chart=["1E2761", "3B6FE8", "7FA6F5", "F5A623", "CADCFC", "5A6483"]),
    "forest": dict(dark="1E4620", primary="2C5F2D", accent="97BC62", on_accent="1D2A1E",
                   accent_text="2C5F2D", ink="1D2A1E", muted="5B6B5C", sub="D6E8B8", bg="FFFFFF",
                   chart=["2C5F2D", "97BC62", "5E9E3A", "B8862B", "C9DDA6", "5B6B5C"]),
    "coral": dict(dark="2F3C7E", primary="2F3C7E", accent="F96167", on_accent="FFFFFF",
                  accent_text="C93A41", ink="20233A", muted="5F6480", sub="F9E795", bg="FFFFFF",
                  chart=["2F3C7E", "F96167", "F9E795", "8A93C9", "C93A41", "5F6480"]),
    "terracotta": dict(dark="5A2A22", primary="B85042", accent="A7BEAE", on_accent="2E211E",
                       accent_text="9C3F33", ink="2E211E", muted="7A6862", sub="E7E8D1", bg="FFFFFF",
                       chart=["B85042", "A7BEAE", "E0A96D", "7A382E", "6F8F7C", "7A6862"]),
    "teal": dict(dark="064E5A", primary="028090", accent="02C39A", on_accent="0B2B30",
                 accent_text="017060", ink="12313A", muted="55737A", sub="BFEDE6", bg="FFFFFF",
                 chart=["028090", "02C39A", "7FD8C8", "F2A541", "064E5A", "55737A"]),
    "charcoal": dict(dark="232D33", primary="36454F", accent="E07A2F", on_accent="FFFFFF",
                     accent_text="B85A1A", ink="212121", muted="5F6B73", sub="D5DBDF", bg="FFFFFF",
                     chart=["36454F", "E07A2F", "8FA3AF", "C5CDD2", "B85A1A", "5F6B73"]),
}

HEAD_FONT, BODY_FONT = "Calibri", "Calibri"


def _effects(shape: Any, shadow: bool) -> None:
    """Force subtle soft shadow or clear effects."""
    sp_pr = shape._element.spPr
    for e in sp_pr.findall(qn("a:effectLst")):
        sp_pr.remove(e)
    eff = etree.SubElement(sp_pr, qn("a:effectLst"))
    if shadow:
        sh = etree.SubElement(eff, qn("a:outerShdw"), blurRad="127000", dist="38100",
                              dir="5400000", algn="t", rotWithShape="0")
        clr = etree.SubElement(sh, qn("a:srgbClr"), val="000000")
        etree.SubElement(clr, qn("a:alpha"), val="14000")


# --------------------------------------------------------------------------- #
# Deck Design SDK
# --------------------------------------------------------------------------- #
class Deck:
    """Flexible PowerPoint presentation builder giving the LLM full composition control."""

    def __init__(
        self,
        theme: str = "midnight",
        overrides: dict[str, Any] | None = None,
        head_font: str = HEAD_FONT,
        body_font: str = BODY_FONT,
    ) -> None:
        t = dict(THEMES.get(theme, THEMES["midnight"]))
        t.update(overrides or {})
        t.setdefault("bg", "FFFFFF")
        t.setdefault("dark2", _mix(t["dark"], "FFFFFF", 0.10))
        t.setdefault("surface", _mix(t["primary"], "FFFFFF", 0.94))
        t.setdefault("line", _mix(t["muted"], "FFFFFF", 0.65))
        self.t = t
        self.theme_name = theme
        self.head, self.body = head_font, body_font
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(SW), Inches(SH)
        self.warnings: list[str] = []
        self._n = 0
        self._num: dict[int, int] = {}

    def color(self, name_or_hex: str) -> str:
        """Resolve a color name from current theme or return clean hex."""
        val = self.t.get(name_or_hex, name_or_hex)
        return str(val).lstrip("#")

    # ---- Core Slide Creation ---------------------------------------------- #
    def add_slide(self, bg: str | None = None) -> Any:
        """Create a blank slide with custom background color (hex or theme key)."""
        bg_hex = self.color(bg or self.t["bg"])
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])  # Blank layout
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = _rgb(bg_hex)
        self._n += 1
        self._num[s.slide_id] = self._n
        return s

    def _warn(self, s: Any, msg: str) -> None:
        self.warnings.append(f"slide {self._num.get(s.slide_id, '?')}: {msg}")

    def _bounds(self, s: Any, x: float, y: float, w: float, h: float, what: str) -> None:
        if x < -0.01 or y < -0.01 or x + w > SW + 0.05 or y + h > SH + 0.05:
            self._warn(s, f"{what} extends past the slide edge: ({round(x,2)}, {round(y,2)}, {round(w,2)}, {round(h,2)})")

    # ---- Primitive: Text Box ---------------------------------------------- #
    def textbox(
        self,
        slide: Any,
        text: str,
        box: tuple[float, float, float, float] = BODY,
        font: str | None = None,
        size: float = 14,
        bold: bool = False,
        italic: bool = False,
        color: str | None = None,
        align: str = "left",           # left | center | right
        vertical_align: str = "top",   # top | middle | bottom
        margin: float = 0,
        wrap: bool = True,
        line_spacing: float | None = None,
        space_after: float = 0,
        fit: bool = True,
        min_pt: float = 10,
        warn: bool = True,
    ) -> float:
        """Add a text box with precise typography and alignment control."""
        x, y, w, h = box
        font = font or self.body
        hex_color = self.color(color or self.t["ink"])
        text = str(text)
        self._bounds(slide, x, y, w, h, "text box")

        if fit:
            size = fit_pt(text, w, h, size, min_pt, font, bold, space_after)
        if warn and est_height(text, w, size, font, bold, space_after) > h * 1.08:
            self._warn(slide, f"text may overflow box: '{text[:35]}...'")

        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = wrap
        tf.margin_left = Inches(margin)
        tf.margin_right = Inches(margin)
        tf.margin_top = Inches(margin)
        tf.margin_bottom = Inches(margin)

        v_map = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM,
                 "t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}
        tf.vertical_anchor = v_map.get(vertical_align, MSO_ANCHOR.TOP)

        a_map = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT,
                 "l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
        align_choice = a_map.get(align, PP_ALIGN.LEFT)

        for i, para in enumerate(text.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align_choice
            if space_after:
                p.space_after = Pt(space_after)
            if line_spacing:
                p.line_spacing = line_spacing
            r = p.add_run()
            r.text = para
            f = r.font
            f.size, f.bold, f.italic, f.name = Pt(size), bold, italic, font
            f.color.rgb = _rgb(hex_color)
        return size

    # Backward compatibility alias
    _text = textbox

    # ---- Primitive: Shapes ------------------------------------------------ #
    def shape(
        self,
        slide: Any,
        kind: str = "rounded_rect",     # rect | rounded_rect | circle | oval
        box: tuple[float, float, float, float] = BODY,
        fill: str | None = None,
        border: str | None = None,
        border_width: float = 1.0,
        radius: float | None = 0.05,
        shadow: bool = False,
        decor: bool = False,
    ) -> Any:
        """Add a geometric shape with fill and border styling."""
        x, y, w, h = box
        if not decor:
            self._bounds(slide, x, y, w, h, "shape")

        kind_map = {
            "rect": MSO_SHAPE.RECTANGLE,
            "rectangle": MSO_SHAPE.RECTANGLE,
            "rounded_rect": MSO_SHAPE.ROUNDED_RECTANGLE,
            "rounded_rectangle": MSO_SHAPE.ROUNDED_RECTANGLE,
            "circle": MSO_SHAPE.OVAL,
            "oval": MSO_SHAPE.OVAL,
        }
        mso_kind = kind_map.get(kind.lower(), MSO_SHAPE.ROUNDED_RECTANGLE)

        shp = slide.shapes.add_shape(mso_kind, Inches(x), Inches(y), Inches(w), Inches(h))
        if fill:
            shp.fill.solid()
            shp.fill.fore_color.rgb = _rgb(self.color(fill))
        else:
            shp.fill.background()

        if border:
            shp.line.color.rgb = _rgb(self.color(border))
            shp.line.width = Pt(border_width)
        else:
            shp.line.fill.background()

        if radius is not None and mso_kind == MSO_SHAPE.ROUNDED_RECTANGLE:
            shp.adjustments[0] = radius

        _effects(shp, shadow)
        return shp

    # Backward compatibility alias
    _shape = shape

    # ---- Primitive: Lines & Connectors ------------------------------------ #
    def line(
        self,
        slide: Any,
        start: tuple[float, float],
        end: tuple[float, float],
        color: str | None = None,
        width: float = 1.0,
        dashed: bool = False,
    ) -> Any:
        """Draw a connecting line or visual divider."""
        x1, y1 = start
        x2, y2 = end
        connector = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
        )
        line_color = self.color(color or self.t["line"])
        connector.line.color.rgb = _rgb(line_color)
        connector.line.width = Pt(width)
        return connector

    # ---- Primitive: Tables ------------------------------------------------ #
    def table(
        self,
        slide: Any,
        rows: Sequence[Sequence[Any]],
        box: tuple[float, float, float, float] = BODY,
        header: bool = True,
        col_widths: Sequence[float] | None = None,
        header_bg: str | None = None,
        header_color: str = "FFFFFF",
        alt_bg: str | None = None,
        text_color: str | None = None,
        font_size: float = 13,
        border_color: str | None = None,
    ) -> Any:
        """Add a professionally styled data table."""
        x, y, w, h = box
        self._bounds(slide, x, y, w, h, "table")
        num_rows = len(rows)
        if num_rows == 0:
            return None
        num_cols = len(rows[0])

        tbl_shape = slide.shapes.add_table(num_rows, num_cols, Inches(x), Inches(y), Inches(w), Inches(h))
        tbl = tbl_shape.table

        # Column widths
        if col_widths and len(col_widths) == num_cols:
            for c_idx, cw in enumerate(col_widths):
                tbl.columns[c_idx].width = Inches(cw)

        h_bg = self.color(header_bg or self.t["primary"])
        h_col = self.color(header_color)
        b_text = self.color(text_color or self.t["ink"])
        alt_hex = self.color(alt_bg or self.t["surface"]) if alt_bg else None

        for r_idx, row in enumerate(rows):
            is_header = header and r_idx == 0
            for c_idx, val in enumerate(row):
                cell = tbl.cell(r_idx, c_idx)
                cell.text = str(val)
                cell.fill.solid()
                if is_header:
                    cell.fill.fore_color.rgb = _rgb(h_bg)
                elif alt_hex and r_idx % 2 == 1:
                    cell.fill.fore_color.rgb = _rgb(alt_hex)
                else:
                    cell.fill.fore_color.rgb = _rgb("FFFFFF")

                tf = cell.text_frame
                tf.word_wrap = True
                tf.margin_left = Inches(0.12)
                tf.margin_right = Inches(0.12)
                tf.margin_top = Inches(0.08)
                tf.margin_bottom = Inches(0.08)
                p = tf.paragraphs[0]
                p.alignment = PP_ALIGN.LEFT
                if len(p.runs) > 0:
                    f = p.runs[0].font
                    f.name = self.body
                    f.size = Pt(font_size + (1 if is_header else 0))
                    f.bold = is_header
                    f.color.rgb = _rgb(h_col if is_header else b_text)

        return tbl_shape

    # ---- Primitive: Images ------------------------------------------------ #
    def image(
        self,
        slide: Any,
        path: str | Path,
        box: tuple[float, float, float, float] = BODY,
    ) -> Any:
        """Place an image onto the slide at the specified box coordinates."""
        img_path = Path(path)
        if not img_path.exists():
            self._warn(slide, f"image not found: {path}")
            return None
        x, y, w, h = box
        self._bounds(slide, x, y, w, h, "image")
        return slide.shapes.add_picture(str(img_path), Inches(x), Inches(y), Inches(w), Inches(h))

    # ---- Primitive: Charts ------------------------------------------------ #
    def chart(
        self,
        slide: Any,
        kind: str,
        categories: Sequence[Any],
        series: dict[str, Sequence[Any]],
        box: tuple[float, float, float, float] = BODY,
        title: str | None = None,
        labels: bool | None = None,
        legend: bool | None = None,
        fmt: str = "General",
        colors: list[str] | None = None,
    ) -> Any:
        """Render a chart (column | bar | stacked | line | area | pie | doughnut)."""
        t = self.t
        x, y, w, h = box
        self._bounds(slide, x, y, w, h, "chart")
        types = {
            "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
            "bar": XL_CHART_TYPE.BAR_CLUSTERED,
            "stacked": XL_CHART_TYPE.COLUMN_STACKED,
            "line": XL_CHART_TYPE.LINE_MARKERS,
            "area": XL_CHART_TYPE.AREA,
            "pie": XL_CHART_TYPE.PIE,
            "doughnut": XL_CHART_TYPE.DOUGHNUT,
        }
        cd = CategoryChartData()
        cd.categories = list(categories)
        for name, vals in series.items():
            cd.add_series(name, list(vals))

        ch = slide.shapes.add_chart(types[kind], Inches(x), Inches(y), Inches(w), Inches(h), cd).chart
        ch.font.size, ch.font.name = Pt(12), self.body
        ch.font.color.rgb = _rgb(t["muted"])
        ch.has_title = bool(title)
        if title:
            ch.chart_title.text_frame.text = title
            p0 = ch.chart_title.text_frame.paragraphs[0]
            if p0.runs:
                f = p0.runs[0].font
                f.size, f.bold, f.name = Pt(15), True, self.body
                f.color.rgb = _rgb(t["ink"])
            ch.chart_title.include_in_layout = False

        pie = kind in ("pie", "doughnut")
        ch.has_legend = legend if legend is not None else (len(series) > 1 or pie)
        if ch.has_legend:
            ch.legend.position = XL_LEGEND_POSITION.BOTTOM
            ch.legend.include_in_layout = False
            ch.legend.font.size = Pt(12)

        pal = colors or t["chart"]
        plot = ch.plots[0]
        if pie:
            plot.vary_by_categories = True
            for i in range(len(categories)):
                pt_point = plot.series[0].points[i]
                c = _rgb(self.color(pal[i % len(pal)]))
                pt_point.format.fill.solid()
                pt_point.format.fill.fore_color.rgb = c
        else:
            for i, ser in enumerate(plot.series):
                c = _rgb(self.color(pal[i % len(pal)]))
                if kind == "line":
                    ser.format.line.color.rgb = c
                    ser.format.line.width = Pt(2.5)
                    ser.marker.style = XL_MARKER_STYLE.CIRCLE
                    ser.marker.size = 7
                    ser.marker.format.fill.solid()
                    ser.marker.format.fill.fore_color.rgb = c
                    ser.marker.format.line.color.rgb = c
                else:
                    ser.format.fill.solid()
                    ser.format.fill.fore_color.rgb = c
            if kind in ("column", "bar", "stacked"):
                plot.gap_width = 70
            if kind == "bar":
                ch.category_axis.reverse_order = True
            ch.category_axis.has_major_gridlines = False
            ch.category_axis.format.line.color.rgb = _rgb(t["line"])
            ch.category_axis.tick_labels.font.size = Pt(12)
            ch.value_axis.tick_labels.font.size = Pt(11)
            ch.value_axis.format.line.fill.background()
            ch.value_axis.major_gridlines.format.line.color.rgb = _rgb(_mix(t["line"], "FFFFFF", 0.5))

        if labels is None:
            labels = kind != "stacked"
        if labels:
            plot.has_data_labels = True
            dl = plot.data_labels
            dl.font.size, dl.font.bold = Pt(12), True
            dl.font.color.rgb = _rgb(t["ink"])
            dl.number_format, dl.number_format_is_linked = fmt, False
            if kind in ("column", "bar"):
                dl.position = XL_LABEL_POSITION.OUTSIDE_END
            elif kind == "line":
                dl.position = XL_LABEL_POSITION.ABOVE
            elif kind == "pie":
                dl.position = XL_LABEL_POSITION.OUTSIDE_END
            if not pie and len(series) == 1 and kind in ("column", "bar"):
                ch.value_axis.visible = False
                ch.value_axis.has_major_gridlines = False
        return ch

    # ---- Primitive: Badge ------------------------------------------------- #
    def badge(
        self,
        slide: Any,
        x: float,
        y: float,
        d: float,
        label: Any,
        fill: str | None = None,
        color: str | None = None,
        size: float | None = None,
    ) -> Any:
        """Add a circular chip or sequence indicator."""
        shp = self.shape(slide, "circle", (x, y, d, d), fill=fill or self.t["accent"])
        tf = shp.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.word_wrap = False
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = str(label)
        r.font.size = Pt(size or max(10, d * 72 * 0.42))
        r.font.bold = True
        r.font.name = self.head
        r.font.color.rgb = _rgb(self.color(color or self.t["on_accent"]))
        return shp

    _badge = badge

    # ---- Composable High-Level Blocks (Optional Convenience) -------------- #
    def notes(self, slide: Any, text: str | None) -> None:
        """Attach speaker notes to slide."""
        if text:
            slide.notes_slide.notes_text_frame.text = text

    def title_slide(self, title: str, subtitle: str = "", notes: str | None = None) -> Any:
        """Convenience starter: Dark theme title slide."""
        t = self.t
        s = self.add_slide(t["dark"])
        self.shape(s, "circle", (8.0, 1.2, 5.0, 5.0), fill=t["dark2"])
        self.shape(s, "circle", (7.1, 4.9, 1.5, 1.5), fill=t["accent"])
        self.textbox(s, title, box=(M, 1.9, 6.8, 2.6), size=44, color="FFFFFF", font=self.head, bold=True, vertical_align="bottom", min_pt=28)
        if subtitle:
            self.textbox(s, subtitle, box=(M, 4.75, 6.8, 1.2), size=20, color=t["sub"], min_pt=14)
        self.notes(s, notes)
        return s

    def content_slide(self, title: str, kicker: str | None = None, notes: str | None = None) -> Any:
        """Convenience starter: Light slide with header and footer."""
        t = self.t
        s = self.add_slide(t["bg"])
        ty = 0.55
        if kicker:
            self.textbox(s, kicker.upper(), box=(M, 0.45, CW, 0.3), size=12, color=t["accent_text"], bold=True, fit=False)
            ty = 0.8
        self.textbox(s, title, box=(M, ty, CW, 0.9), size=32, color=t["primary"], font=self.head, bold=True, min_pt=24)
        self.textbox(s, str(self._n), box=(SW - M - 1.0, 7.0, 1.0, 0.25), size=11, color=t["muted"], align="right", fit=False)
        self.notes(s, notes)
        return s

    def section_slide(self, title: str, number: int | None = None, notes: str | None = None) -> Any:
        """Convenience starter: Dark divider section slide."""
        t = self.t
        s = self.add_slide(t["dark"])
        if number is not None:
            self.textbox(s, str(number).zfill(2), box=(M, 1.7, 5.0, 1.7), size=96, color=t["accent"], font=self.head, bold=True, fit=False)
        self.textbox(s, title, box=(M, 3.6, 10.5, 1.8), size=42, color="FFFFFF", font=self.head, bold=True, min_pt=28)
        self.notes(s, notes)
        return s

    def quote_slide(self, quote: str, who: str = "", notes: str | None = None) -> Any:
        """Convenience starter: Impactful quote slide."""
        t = self.t
        s = self.add_slide(t["surface"])
        self.textbox(s, "\u201C", box=(M, 0.9, 2.0, 1.6), size=140, color=t["accent"], font=self.head, bold=True, fit=False)
        self.textbox(s, quote, box=(M + 0.3, 2.4, CW - 1.2, 2.9), size=32, color=t["primary"], font=self.head, italic=True, vertical_align="bottom", min_pt=22)
        if who:
            self.textbox(s, who, box=(M + 0.3, 5.6, CW - 1.2, 0.5), size=18, color=t["muted"], bold=True)
        self.notes(s, notes)
        return s

    def closing_slide(self, title: str, lines: list[str] | None = None, notes: str | None = None) -> Any:
        """Convenience starter: Clean closing/next-steps slide."""
        t = self.t
        s = self.add_slide(t["dark"])
        self.shape(s, "circle", (9.6, 3.9, 3.6, 3.6), fill=t["dark2"])
        self.textbox(s, title, box=(M, 1.0, 10.5, 1.4), size=40, color="FFFFFF", font=self.head, bold=True, min_pt=28)
        if lines:
            self.textbox(s, "\n".join(lines), box=(M, 2.8, 8.0, 3.6), size=22, color=t["sub"], space_after=12, min_pt=14)
        self.notes(s, notes)
        return s

    def bullets(
        self,
        slide: Any,
        items: list[str],
        box: tuple[float, float, float, float] = BODY,
        size: float = 18,
        color: str | None = None,
    ) -> None:
        """Render a formatted bullet point list."""
        x, y, w, h = box
        color_hex = self.color(color or self.t["ink"])
        gap = round(size * 0.5)
        size = fit_pt("\n".join(items), w - 0.35, h, size, 13, self.body, False, gap)
        self._bounds(slide, x, y, w, h, "bullet list")
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for i, it in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(gap)
            r = p.add_run()
            r.text = str(it)
            r.font.size, r.font.name = Pt(size), self.body
            r.font.color.rgb = _rgb(color_hex)
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(Inches(0.35)))
            pPr.set("indent", str(-Inches(0.35)))
            clr = etree.SubElement(pPr, qn("a:buClr"))
            etree.SubElement(clr, qn("a:srgbClr"), val=self.t["accent"])
            etree.SubElement(pPr, qn("a:buFont"), typeface="Arial")
            etree.SubElement(pPr, qn("a:buChar"), char="\u2022")

    def cards(
        self,
        slide: Any,
        items: list[dict[str, Any]],
        cols: int | None = None,
        box: tuple[float, float, float, float] = BODY,
        badges: bool = True,
    ) -> None:
        """Auto-arranged container cards: items=[{'title', 'body', 'badge'}]."""
        t, n = self.t, len(items)
        if n == 0:
            return
        cols = cols or (n if n <= 3 else math.ceil(n / 2))
        rows = math.ceil(n / cols)
        pad_size = 0.3
        for i, (it, (x, y, w, h)) in enumerate(zip(items, grid(box, cols, rows))):
            self.shape(slide, "rounded_rect", (x, y, w, h), fill=t["surface"], radius=0.05, shadow=True)
            label = it.get("badge", i + 1)
            if cols == 1:
                d = 0.7
                if badges:
                    self.badge(slide, x + pad_size, y + (h - d) / 2, d, label)
                tx = x + pad_size + (d + 0.3 if badges else 0)
                tw = x + w - pad_size - tx
                self.textbox(slide, it["title"], box=(tx, y + pad_size, tw, 0.45), size=20, color=t["primary"], font=self.head, bold=True, min_pt=15)
                self.textbox(slide, it.get("body", ""), box=(tx, y + pad_size + 0.5, tw, h - 2 * pad_size - 0.5), size=15, min_pt=11)
            elif h < 3.2:
                d = 0.6
                tx = x + pad_size + (d + 0.25 if badges else 0)
                if badges:
                    self.badge(slide, x + pad_size, y + pad_size, d, label)
                self.textbox(slide, it["title"], box=(tx, y + pad_size, x + w - pad_size - tx, d), size=19, color=t["primary"], font=self.head, bold=True, vertical_align="middle", min_pt=14)
                by = y + pad_size + d + 0.2
                self.textbox(slide, it.get("body", ""), box=(x + pad_size, by, w - 2 * pad_size, y + h - pad_size - by), size=15, min_pt=11)
            else:
                ty = y + pad_size
                if badges:
                    self.badge(slide, x + pad_size, ty, 0.65, label)
                    ty += 0.65 + 0.2
                self.textbox(slide, it["title"], box=(x + pad_size, ty, w - 2 * pad_size, 0.7), size=19, color=t["primary"], font=self.head, bold=True, min_pt=14)
                by = ty + 0.75
                self.textbox(slide, it.get("body", ""), box=(x + pad_size, by, w - 2 * pad_size, y + h - pad_size - by), size=15, min_pt=11)

    def kpis(
        self,
        slide: Any,
        items: list[dict[str, Any]],
        box: tuple[float, float, float, float] = (M, 2.1, CW, 2.8),
    ) -> None:
        """Key Performance Indicator cards in a row: items=[{'value', 'label', 'delta'}]."""
        t = self.t
        pad_size = 0.25
        for it, (x, y, w, h) in zip(items, grid(box, len(items), 1)):
            self.shape(slide, "rounded_rect", (x, y, w, h), fill=t["surface"], radius=0.05, shadow=True)
            self.textbox(slide, it["value"], box=(x + pad_size, y + pad_size, w - 2 * pad_size, 1.1), size=44, color=t["primary"], font=self.head, bold=True, vertical_align="middle", min_pt=24)
            delta = it.get("delta")
            lab_h = h - 1.3 - pad_size - (0.4 if delta else 0)
            self.textbox(slide, it["label"], box=(x + pad_size, y + pad_size + 1.15, w - 2 * pad_size, lab_h), size=16, min_pt=12)
            if delta:
                col = "2E7D32" if str(delta).lstrip().startswith("+") else (
                    "C62828" if str(delta).lstrip().startswith(("-", "\u2212")) else t["muted"])
                self.textbox(slide, delta, box=(x + pad_size, y + h - pad_size - 0.35, w - 2 * pad_size, 0.35), size=14, color=col, bold=True, fit=False)

    def callout(
        self,
        slide: Any,
        text: str,
        box: tuple[float, float, float, float],
        label: str | None = None,
    ) -> None:
        """Emphasized highlight block."""
        t = self.t
        x, y, w, h = box
        self.shape(slide, "rounded_rect", (x, y, w, h), fill=t["primary"], radius=0.05, shadow=True)
        ty = y + 0.4
        if label:
            self.textbox(slide, label.upper(), box=(x + 0.4, ty, w - 0.8, 0.3), size=12, color=t["sub"], bold=True, fit=False)
            ty += 0.45
        self.textbox(slide, text, box=(x + 0.4, ty, w - 0.8, y + h - 0.4 - ty), size=24, color="FFFFFF", font=self.head, bold=True, vertical_align="middle", min_pt=15)

    def timeline(
        self,
        slide: Any,
        steps: list[dict[str, Any]],
        box: tuple[float, float, float, float] | None = None,
        orient: str = "h",
    ) -> None:
        """Visual roadmap or process timeline."""
        t, n = self.t, len(steps)
        if n == 0:
            return
        if orient == "h":
            x, y, w, h = box or (M, 2.4, CW, 4.2)
            colw, d = w / n, 0.75
            self.shape(slide, "rect", (x + d / 2, y + d / 2 - 0.02, (n - 1) * colw, 0.04), fill=t["line"])
            for i, st in enumerate(steps):
                cx = x + i * colw
                self.badge(slide, cx, y, d, i + 1, fill=t["primary"], color="FFFFFF")
                self.textbox(slide, st["title"], box=(cx, y + d + 0.3, colw - 0.35, 0.55), size=18, color=t["primary"], font=self.head, bold=True, min_pt=14)
                self.textbox(slide, st.get("body", ""), box=(cx, y + d + 0.95, colw - 0.35, h - d - 0.95), size=14, min_pt=11)
        else:
            x, y, w, h = box or BODY
            rowh, d = h / n, 0.6
            self.shape(slide, "rect", (x + d / 2 - 0.02, y + d / 2, 0.04, (n - 1) * rowh), fill=t["line"])
            for i, st in enumerate(steps):
                ry = y + i * rowh
                self.badge(slide, x, ry, d, i + 1, fill=t["primary"], color="FFFFFF")
                self.textbox(slide, st["title"], box=(x + d + 0.3, ry, w - d - 0.3, 0.4), size=18, color=t["primary"], font=self.head, bold=True, min_pt=14)
                self.textbox(slide, st.get("body", ""), box=(x + d + 0.3, ry + 0.42, w - d - 0.3, max(0.3, rowh - 0.55)), size=14, min_pt=11)

    def two_col(
        self,
        slide: Any,
        left: dict[str, Any],
        right: dict[str, Any],
        box: tuple[float, float, float, float] = BODY,
    ) -> None:
        """Side-by-side comparison panels."""
        t = self.t
        for side, (x, y, w, h) in zip((left, right), split(box, [1, 1])):
            self.shape(slide, "rounded_rect", (x, y, w, h), fill=t["surface"], radius=0.04, shadow=True)
            self.textbox(slide, side["heading"], box=(x + 0.4, y + 0.35, w - 0.8, 0.6), size=22, color=t["primary"], font=self.head, bold=True, min_pt=16)
            self.bullets(slide, side["bullets"], box=(x + 0.4, y + 1.15, w - 0.8, h - 1.5), size=18)

    # ---- Save Contract ---------------------------------------------------- #
    def save(self, path: str = "deck.pptx") -> dict[str, Any]:
        """Save and return standard metadata dict for the workbench agent."""
        self.prs.save(path)
        return {"path": path, "slides": self._n, "warnings": list(self.warnings)}
