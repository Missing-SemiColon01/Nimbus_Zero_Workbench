from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

class DocxBuilder:
    def __init__(self, theme="modern"):
        self.doc = Document()
        self.theme = theme
        self._setup_styles()

    def _setup_styles(self):
        style = self.doc.styles['Normal']
        font = style.font
        font.name = 'Arial'
        font.size = Pt(11)
        font.color.rgb = RGBColor(0x33, 0x33, 0x33)

        title_style = self.doc.styles['Title']
        title_font = title_style.font
        title_font.name = 'Arial'
        title_font.size = Pt(24)
        title_font.bold = True
        title_font.color.rgb = RGBColor(0x0f, 0x17, 0x2a)  # Navy

    def add_title(self, text, subtitle=None):
        self.doc.add_heading(text, 0)
        if subtitle:
            p = self.doc.add_paragraph(subtitle)
            p.style.font.size = Pt(14)
            p.style.font.color.rgb = RGBColor(0x64, 0x74, 0x8b)

    def add_heading(self, text, level=1):
        h = self.doc.add_heading(text, level=level)
        h.style.font.color.rgb = RGBColor(0x0f, 0x17, 0x2a)

    def add_paragraph(self, text):
        return self.doc.add_paragraph(text)

    def add_bullet(self, text):
        return self.doc.add_paragraph(text, style='List Bullet')

    def save(self, path):
        self.doc.save(path)
        return {"status": "success", "file": str(path)}
